

from __future__ import annotations

import re
import logging
from datetime import datetime
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Optional

import pdfplumber
from openpyxl import load_workbook

# OCR opsiyonel
try:
    import pytesseract
    from pdf2image import convert_from_path
    OCR_KULLANILABILIR = True
except ImportError:
    OCR_KULLANILABILIR = False

# LLM opsiyonel - import edilemese bile motor çalışmalı
try:
    from llm_cikarim import (
        llm_ile_cikar,
        regex_ve_llm_birlestir,
    )
    from llm_saglayici import LLMSaglayici
    LLM_KULLANILABILIR = True
except ImportError:
    LLM_KULLANILABILIR = False
    LLMSaglayici = None  # type: ignore


logger = logging.getLogger("ihale_otomasyonu")
logger.setLevel(logging.INFO)
if not logger.handlers:
    handler = logging.StreamHandler()
    handler.setFormatter(
        logging.Formatter("%(asctime)s [%(levelname)s] %(message)s",
                          datefmt="%Y-%m-%d %H:%M:%S")
    )
    logger.addHandler(handler)


# ============================================================
# VERİ YAPISI
# ============================================================
@dataclass
class IhaleVerisi:
    ihale_kayit_no: str = ""
    ihale_konusu: str = ""
    idare_adi: str = ""
    ihale_usulu: str = ""
    ihale_tarihi: str = ""
    son_teklif_tarihi: str = ""
    yaklasik_maliyet: str = ""
    para_birimi: str = "TL"
    ihale_yeri: str = ""
    iletisim: str = ""
    pdf_dosya_adi: str = ""

    def bos_alan_sayisi(self) -> int:
        return sum(1 for k, v in asdict(self).items()
                   if k not in ("pdf_dosya_adi", "para_birimi") and not v)


# ============================================================
# PDF OKUMA
# ============================================================
def pdf_metin_oku(pdf_yolu: Path) -> str:
    """PDF'ten metin okur. Metin tabanlı PDF için pdfplumber, taranmış için OCR."""
    metin = ""
    try:
        with pdfplumber.open(pdf_yolu) as pdf:
            for sayfa in pdf.pages:
                metin += (sayfa.extract_text() or "") + "\n"
    except Exception as e:
        logger.warning(f"pdfplumber hatası: {e}")

    if len(metin.strip()) < 100:
        if not OCR_KULLANILABILIR:
            logger.warning(
                "PDF taranmış görünüyor ama OCR kütüphaneleri yok. "
                "LLM aktifse PDF doğrudan LLM'e gönderilecek."
            )
            return metin
        logger.info(f"PDF taranmış, OCR çalışıyor: {pdf_yolu.name}")
        try:
            metin = ocr_ile_oku(pdf_yolu)
        except Exception as e:
            logger.warning(f"OCR hatası: {e}")
    return metin


def ocr_ile_oku(pdf_yolu: Path) -> str:
    metin = ""
    sayfalar = convert_from_path(str(pdf_yolu), dpi=300)
    for i, sayfa in enumerate(sayfalar, 1):
        try:
            sayfa_metni = pytesseract.image_to_string(sayfa, lang="tur+eng")
        except pytesseract.TesseractError:
            sayfa_metni = pytesseract.image_to_string(sayfa, lang="eng")
        metin += f"\n--- Sayfa {i} ---\n{sayfa_metni}\n"
    return metin


# ============================================================
# REGEX DESENLERİ
# ============================================================
DESENLER = {
    "ihale_kayit_no": [
        r"İhale\s*Kay[ıi]t\s*Numaras[ıi]\s*[:\-]?\s*([0-9]{4}\s*/\s*[0-9]{3,8})",
        r"İKN\s*[:\-]?\s*([0-9]{4}\s*/\s*[0-9]{3,8})",
        r"İhale\s*No(?:su)?\s*[:\-]?\s*([0-9]{4}\s*/\s*[0-9]{3,8})",
    ],
    "ihale_konusu": [
        r"İhale(?:nin)?\s*Konusu\s*[:\-]\s*(.+?)(?=\n)",
        r"İşin\s*(?:Ad[ıi]|Konusu|Niteli[ğg]i)\s*[:\-]\s*(.+?)(?=\n)",
    ],
    "idare_adi": [
        r"İdarenin\s*\n[^\n]*?a\s*\)\s*Ad[ıi]\s*[:\-]\s*(.+?)(?=\n)",
        r"a\s*\)\s*Ad[ıi]\s*[:\-]\s*(.+?)(?=\n)",
        r"İdare(?:nin)?\s*Ad[ıi]\s*[:\-]\s*(.+?)(?=\n)",
        r"İhaleyi\s*Yapan\s*İdare\s*[:\-]\s*(.+?)(?=\n)",
    ],
    "ihale_usulu": [
        r"a\s*\)\s*İhale\s*Usul[üu]\s*[:\-]\s*(.+?)(?=\n)",
        r"İhale\s*Usul[üu]\s*[:\-]\s*(.+?)(?=\n)",
        r"İhalenin\s*Usul[üu]\s*[:\-]\s*(.+?)(?=\n)",
        r"\b(A[çc][ıi]k\s*İhale\s*Usul[üu])\b",
        r"\b(Pazarl[ıi]k\s*Usul[üu])\b",
        r"\b(Belli\s*İstekliler\s*Aras[ıi]nda(?:\s*İhale\s*Usul[üu])?)\b",
        r"\b(Do[ğg]rudan\s*Temin)\b",
    ],
    "ihale_tarihi": [
        r"b\s*\)\s*İhale\s*Tarihi\s*ve\s*Saati\s*[:\-]\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})",
        r"İhale\s*Tarihi\s*(?:ve\s*Saati)?\s*[:\-]\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})",
        r"İhalenin\s*Yap[ıi]laca[ğg][ıi]\s*Tarih\s*[:\-]\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})",
    ],
    "son_teklif_tarihi": [
        r"Son\s*Teklif\s*(?:Verme\s*)?Tarihi\s*(?:ve\s*Saati)?\s*[:\-]\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})",
        r"Tekliflerin\s*Sunulaca[ğg][ıi]\s*(?:Son\s*)?Tarih\s*[:\-]\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})",
        r"Teklif\s*Verme\s*Son\s*Tarihi\s*[:\-]\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})",
    ],
    "yaklasik_maliyet": [
        r"Yakla[şs][ıi]k\s*Maliyet\s*[:\-]?\s*([\d.,]+)\s*(?:TL|₺|TRY|EUR|USD)?",
        r"Tahmini\s*Bedel\s*[:\-]?\s*([\d.,]+)\s*(?:TL|₺|TRY|EUR|USD)?",
        r"Muhammen\s*Bedel\s*[:\-]?\s*([\d.,]+)\s*(?:TL|₺|TRY|EUR|USD)?",
    ],
    "ihale_yeri": [
        r"İhalenin\s*Yap[ıi]laca[ğg][ıi]\s*Yer(?:i)?\s*[:\-]\s*(.+?)(?=\n)",
        r"[çc]\s*\)\s*İhalenin\s*Yap[ıi]laca[ğg][ıi]\s*Yer\s*[:\-]\s*(.+?)(?=\n)",
        r"İhale\s*Yer(?:i)?\s*[:\-]\s*(.+?)(?=\n)",
    ],
    "iletisim": [
        r"Telefon\s*ve\s*Faks[^\n]*?Telefon\s*[:\-]?\s*([\d\s\(\)\+\-/]{10,25})",
        r"(?:Telefon|Tel)\s*(?:No)?\s*[:\-]\s*([\d\s\(\)\+\-/]{10,25})",
        r"(?:E[\-\s]?Posta|E[\-\s]?Mail)\s*[:\-]\s*([\w\.\-]+@[\w\.\-]+\.\w+)",
    ],
}

PARA_DESENI = re.compile(
    r"(?:Yakla[şs][ıi]k\s*Maliyet|Tahmini\s*Bedel|Muhammen\s*Bedel)\s*[:\-]?\s*[\d.,]+\s*(TL|₺|TRY|EUR|USD|€|\$)",
    re.IGNORECASE,
)


def desen_dene(metin: str, desenler: list[str]) -> str:
    for desen in desenler:
        m = re.search(desen, metin, re.IGNORECASE | re.DOTALL)
        if m:
            sonuc = m.group(1).strip()
            sonuc = re.sub(r"\s+", " ", sonuc)
            if len(sonuc) > 200:
                sonuc = sonuc[:200].rsplit(" ", 1)[0] + "..."
            return sonuc
    return ""


# ============================================================
# REGEX TABANLI ÇIKARIM (fallback)
# ============================================================
def _regex_ile_cikar(metin: str, pdf_dosya_adi: str) -> dict:
    sonuc = {
        "ihale_kayit_no": desen_dene(metin, DESENLER["ihale_kayit_no"]),
        "ihale_konusu": desen_dene(metin, DESENLER["ihale_konusu"]),
        "idare_adi": desen_dene(metin, DESENLER["idare_adi"]),
        "ihale_usulu": desen_dene(metin, DESENLER["ihale_usulu"]),
        "ihale_tarihi": desen_dene(metin, DESENLER["ihale_tarihi"]),
        "son_teklif_tarihi": desen_dene(metin, DESENLER["son_teklif_tarihi"]),
        "yaklasik_maliyet": desen_dene(metin, DESENLER["yaklasik_maliyet"]),
        "ihale_yeri": desen_dene(metin, DESENLER["ihale_yeri"]),
        "iletisim": desen_dene(metin, DESENLER["iletisim"]),
        "para_birimi": "TL",
        "pdf_dosya_adi": pdf_dosya_adi,
    }

    if sonuc["ihale_kayit_no"]:
        sonuc["ihale_kayit_no"] = re.sub(r"\s*/\s*", "/", sonuc["ihale_kayit_no"])

    para_match = PARA_DESENI.search(metin)
    if para_match:
        sembol = para_match.group(1).upper()
        sonuc["para_birimi"] = {"₺": "TL", "TRY": "TL", "€": "EUR", "$": "USD"}.get(sembol, sembol)

    return sonuc


# ============================================================
# ANA ÇIKARIM - HİBRİT (regex + LLM)
# ============================================================
def ihale_verisi_cikar(
    metin: str,
    pdf_dosya_adi: str,
    saglayici: Optional["LLMSaglayici"] = None,
    pdf_yolu: Optional[Path] = None,
) -> IhaleVerisi:
    """
    İhale ilanından veri çıkarır.

    Args:
        metin: PDF'ten çıkarılmış düz metin (boş olabilir)
        pdf_dosya_adi: PDF dosya adı
        saglayici: LLM sağlayıcı (None ise sadece regex)
        pdf_yolu: Metin yetersizse LLM'e direkt PDF göndermek için

    Returns:
        IhaleVerisi nesnesi
    """
    # 1. Regex çıkarımı (her zaman)
    regex_sonucu = _regex_ile_cikar(metin, pdf_dosya_adi)

    # 2. LLM çıkarımı (varsa)
    if saglayici is not None and LLM_KULLANILABILIR:
        try:
            llm_yaniti = llm_ile_cikar(
                belge_tipi="ihale_ilani",
                belge_metni=metin,
                regex_sonucu=regex_sonucu,
                saglayici=saglayici,
                pdf_yolu=pdf_yolu,
            )
            if llm_yaniti.basarili and llm_yaniti.veri:
                birlesik = regex_ve_llm_birlestir(
                    regex_sonucu=regex_sonucu,
                    llm_sonucu=llm_yaniti.veri,
                    belge_tipi="ihale_ilani",
                )
                return _dict_to_ihale_verisi(birlesik, pdf_dosya_adi)
            else:
                logger.warning(
                    f"LLM başarısız, regex'e dönülüyor: {llm_yaniti.hata}"
                )
        except Exception as e:
            logger.exception(f"LLM çıkarım hatası, regex fallback: {e}")

    return _dict_to_ihale_verisi(regex_sonucu, pdf_dosya_adi)


def _dict_to_ihale_verisi(d: dict, pdf_dosya_adi: str) -> IhaleVerisi:
    return IhaleVerisi(
        ihale_kayit_no=d.get("ihale_kayit_no", "") or "",
        ihale_konusu=d.get("ihale_konusu", "") or "",
        idare_adi=d.get("idare_adi", "") or "",
        ihale_usulu=d.get("ihale_usulu", "") or "",
        ihale_tarihi=d.get("ihale_tarihi", "") or "",
        son_teklif_tarihi=d.get("son_teklif_tarihi", "") or "",
        yaklasik_maliyet=d.get("yaklasik_maliyet", "") or "",
        para_birimi=d.get("para_birimi", "TL") or "TL",
        ihale_yeri=d.get("ihale_yeri", "") or "",
        iletisim=d.get("iletisim", "") or "",
        pdf_dosya_adi=pdf_dosya_adi,
    )


# ============================================================
# EXCEL'E YAZMA
# ============================================================
def ihaleyi_excele_ekle(excel_yolu: Path, veri: IhaleVerisi, notlar: str = "") -> int:
    if not excel_yolu.exists():
        raise FileNotFoundError(f"İhale Excel dosyası bulunamadı: {excel_yolu}")

    wb = load_workbook(excel_yolu)
    ws = wb.active
    yeni_satir = ws.max_row + 1
    sira_no = yeni_satir - 1

    eklenecek = [
        sira_no,
        veri.ihale_kayit_no,
        veri.ihale_konusu,
        veri.idare_adi,
        veri.ihale_usulu,
        veri.ihale_tarihi,
        veri.son_teklif_tarihi,
        veri.yaklasik_maliyet,
        veri.para_birimi,
        veri.ihale_yeri,
        veri.iletisim,
        veri.pdf_dosya_adi,
        datetime.now().strftime("%Y-%m-%d %H:%M"),
        notlar,
    ]

    for col_idx, deger in enumerate(eklenecek, start=1):
        ws.cell(row=yeni_satir, column=col_idx, value=deger)

    wb.save(excel_yolu)
    return yeni_satir


# Geri uyumluluk
excele_ekle = ihaleyi_excele_ekle


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Kullanım: python ihale_motor.py <pdf_yolu> [--llm <claude|gemini>]")
        sys.exit(1)

    pdf_yolu = Path(sys.argv[1])
    saglayici = None

    if "--llm" in sys.argv:
        try:
            from llm_saglayici import saglayici_olustur
            idx = sys.argv.index("--llm")
            saglayici_adi = sys.argv[idx + 1] if len(sys.argv) > idx + 1 else "claude"
            saglayici = saglayici_olustur(saglayici_adi)
            print(f"LLM aktif: {saglayici.ad}/{saglayici.model}")
        except Exception as e:
            print(f"LLM kurulamadı: {e}\nSadece regex ile devam ediliyor...")

    metin = pdf_metin_oku(pdf_yolu)
    veri = ihale_verisi_cikar(metin, pdf_yolu.name, saglayici=saglayici, pdf_yolu=pdf_yolu)

    print("\n=== Çıkarılan İhale Verisi ===")
    for k, v in asdict(veri).items():
        print(f"  {k:25s} : {v}")
    print(f"\nBoş alan sayısı: {veri.bos_alan_sayisi()} / 9")
